"""Eigene Firmenlisten (CSV/XLSX) als Alternative zur Google-Places-Suche einlesen."""

from __future__ import annotations

import csv
from pathlib import Path

from leadscraper.dedupe import normalize_domain
from leadscraper.geo import bundesland_from_plz, normalize_bundesland
from leadscraper.models import Company

# Spaltenname (lowercase, ohne Sonderzeichen) → Company-Feld
_COLUMN_ALIASES: dict[str, str] = {
    "firma": "name",
    "name": "name",
    "unternehmen": "name",
    "firmenname": "name",
    "company": "name",
    "website": "website",
    "webseite": "website",
    "url": "website",
    "homepage": "website",
    "web": "website",
    "telefon": "phone",
    "tel": "phone",
    "phone": "phone",
    "telefonnummer": "phone",
    "festnetz": "phone",
    "strasse": "street",
    "straße": "street",
    "street": "street",
    "adresse": "street",
    "plz": "plz",
    "postleitzahl": "plz",
    "zip": "plz",
    "ort": "city",
    "stadt": "city",
    "city": "city",
    "bundesland": "bundesland",
    "region": "bundesland",
    "state": "bundesland",
    "branche": "primary_type",
    "typ": "primary_type",
    "type": "primary_type",
}


def _norm_header(h: str) -> str:
    return "".join(ch for ch in h.strip().lower() if ch.isalnum() or ch in "äöüß")


def _rows_from_csv(path: Path) -> list[dict[str, str]]:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(raw.splitlines(), dialect=dialect)
    return [{k or "": (v or "") for k, v in row.items()} for row in reader]


def _rows_from_xlsx(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h or "") for h in next(rows_iter, [])]
    rows = []
    for values in rows_iter:
        if not any(v not in (None, "") for v in values):
            continue
        rows.append(
            {header[i]: ("" if v is None else str(v)) for i, v in enumerate(values) if i < len(header)}
        )
    wb.close()
    return rows


def read_company_list(path: Path) -> list[Company]:
    """CSV (; , oder Tab) oder XLSX mit mindestens einer Spalte Firma/Name oder Website."""
    rows = _rows_from_xlsx(path) if path.suffix.lower() in (".xlsx", ".xlsm") else _rows_from_csv(path)
    companies: list[Company] = []
    for i, row in enumerate(rows, start=1):
        data: dict[str, str] = {}
        for col, val in row.items():
            field = _COLUMN_ALIASES.get(_norm_header(col))
            if field and val and val.strip():
                data.setdefault(field, val.strip())
        website = data.get("website")
        if website and "://" not in website:
            website = "https://" + website
        name = data.get("name") or (normalize_domain(website) if website else None)
        if not name:
            continue
        plz = data.get("plz")
        if plz and plz.isdigit() and len(plz) == 4:
            plz = "0" + plz  # Excel frisst führende Nullen
        bundesland = normalize_bundesland(data.get("bundesland")) or bundesland_from_plz(plz)
        companies.append(
            Company(
                place_id=f"list-{i}",
                name=name,
                website=website,
                domain=normalize_domain(website),
                phone=data.get("phone"),
                street=data.get("street"),
                plz=plz,
                city=data.get("city"),
                bundesland=bundesland,
                primary_type=data.get("primary_type"),
                query=f"Liste: {path.name}",
            )
        )
    return companies
