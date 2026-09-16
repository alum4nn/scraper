from pathlib import Path

from openpyxl import load_workbook

from leadscraper.demo import demo_leads
from leadscraper.excel import LEAD_COLUMNS, lead_rows, write_workbook
from leadscraper.funding import funding_reference_rows
from leadscraper.models import SearchSpec

SPEC = SearchSpec(queries=["Immobilienmakler"], city="Köln")


def test_workbook_structure_and_formats(tmp_path: Path):
    leads = demo_leads(SPEC)
    out = write_workbook(leads, SPEC, tmp_path / "demo.xlsx", funding_rows=funding_reference_rows())
    wb = load_workbook(out)
    assert wb.sheetnames == ["Leads", "Entscheider", "Alle Nummern", "Förderung", "Meta"]
    ws = wb["Leads"]
    headers = [c.value for c in ws[1]]
    assert headers == list(LEAD_COLUMNS)
    assert ws.max_row == len(leads) + 1
    assert ws.freeze_panes == "F2"
    assert ws.auto_filter.ref
    col = headers.index("Nummer") + 1
    cell = ws.cell(row=2, column=col)
    assert isinstance(cell.value, str) and cell.number_format == "@"
    web = ws.cell(row=2, column=headers.index("Webseite") + 1)
    assert web.hyperlink is not None and web.hyperlink.target.startswith("https://")
    score_col = headers.index("Score") + 1
    scores = [ws.cell(row=r, column=score_col).value for r in range(2, ws.max_row + 1)]
    assert scores == sorted(scores, reverse=True)
    assert ws.data_validations.dataValidation  # Dropdown "Status Akquise"
    people = wb["Entscheider"]
    first = [c.value for c in people[2]]
    assert first[5]  # erste Zeile hat Handy (Handy-Einträge zuerst)
    prem_col = headers.index("Premium") + 1
    assert [ws.cell(row=r, column=prem_col).value for r in range(2, ws.max_row + 1)].count("ja") == 2
    assert headers[:5] == ["Name", "Unternehmensname", "E-Mail", "Nummer", "Webseite"]
    meta = {row[0].value: row[1].value for row in wb["Meta"].iter_rows(min_row=2)}
    assert meta["Anzahl Firmen"] == len(leads)
    assert meta["Entscheider mit Handy"] == 2
    assert wb["Förderung"].max_row > 5


def test_lead_rows_and_empty_workbook(tmp_path: Path):
    rows = lead_rows(demo_leads(SPEC))
    assert rows[0]["Score"] >= rows[-1]["Score"]
    assert all(set(r) == set(LEAD_COLUMNS) for r in rows)
    out = write_workbook([], SPEC, tmp_path / "empty.xlsx")
    wb = load_workbook(out)
    assert [c.value for c in wb["Leads"][1]] == list(LEAD_COLUMNS)
    assert wb["Förderung"]["A1"].value  # Hinweis statt Tabelle


def test_trello_csv_has_two_columns_and_all_infos(tmp_path: Path):
    import csv

    from leadscraper.excel import write_trello_csv

    leads = demo_leads(SPEC)
    out = write_trello_csv(leads, tmp_path / "trello.csv")
    with open(out, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    assert rows[0] == ["Unternehmensname", "Beschreibung"]
    assert len(rows) == len(leads) + 1
    assert all(len(r) == 2 for r in rows)
    titles = [r[0] for r in rows[1:]]
    assert titles == [ld.display_name for ld in sorted(leads, key=lambda x: -x.score)][: len(titles)]
    body = rows[1][1]
    for label in ("Ansprechpartner", "Handy", "Webseite", "Förderung § 82 SGB III", "Pitch"):
        assert f"**{label}:**" in body
    assert "\n" in body  # mehrzeilige Beschreibung bleibt in einer CSV-Zelle


def test_nummer_faellt_auf_die_zentrale_zurueck():
    """Liste B: Ohne Handynummer gehört die Festnetznummer in die Spalte „Nummer“."""
    from leadscraper.excel import lead_rows
    from leadscraper.models import Company, Enrichment, Lead, Person, PhoneNumber

    lead = Lead(
        company=Company(place_id="z1", name="Beispiel Verwaltung GmbH", phone="+49 221 1234560"),
        enrichment=Enrichment(
            website="https://beispiel.example",
            pages_crawled=["https://beispiel.example/impressum"],
            people=[
                Person(
                    name="Anna Beispiel",
                    role="Geschäftsführerin",
                    role_category="geschaeftsfuehrung",
                )
            ],
            phones=[
                PhoneNumber(
                    raw="0221 1234560",
                    e164="+492211234560",
                    national="0221 1234560",
                    kind="landline",
                    source="impressum",
                )
            ],
        ),
    )
    row = lead_rows([lead])[0]
    assert row["Name"] == "Anna Beispiel"
    assert row["Nummer"] == "0221 1234560"


def test_reihenfolge_des_aufrufers_bleibt_erhalten():
    """Nach Belegschaft sortiert heißt: Das Blatt sortiert nicht heimlich nach Score zurück."""
    from leadscraper import excel
    from leadscraper.excel import lead_rows
    from leadscraper.models import Company, Enrichment, Lead, StaffEvidence

    def lead(name: str, score: int, koepfe: int) -> Lead:
        return Lead(
            company=Company(place_id=name, name=name),
            enrichment=Enrichment(staff=StaffEvidence(headcount=koepfe)),
            score=score,
        )

    leads = [lead("Gross GmbH", 40, 48), lead("Klein GmbH", 90, 21)]
    assert [r["Unternehmensname"] for r in lead_rows(leads)] == ["Klein GmbH", "Gross GmbH"]
    excel.set_sortierung(excel.SORTIERUNG_UEBERNEHMEN)
    try:
        assert [r["Unternehmensname"] for r in lead_rows(leads)] == ["Gross GmbH", "Klein GmbH"]
    finally:
        excel.set_sortierung("score")


def test_steuerzeichen_landen_nicht_im_arbeitsblatt():
    """Ein Steuerzeichen aus einer kaputt kodierten Website hat openpyxl abbrechen lassen –
    und damit das Ergebnis eines ganzen 1.499er-Blocks gekostet."""
    from leadscraper.excel import _cell_value

    assert _cell_value("Kathrin\x05 · 6 99 26") == "Kathrin · 6 99 26"
    assert _cell_value("Zeile\nUmbruch\tTab bleibt") == "Zeile\nUmbruch\tTab bleibt"
    assert _cell_value(42) == 42
