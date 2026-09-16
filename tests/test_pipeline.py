import asyncio
from pathlib import Path

import pytest
from openpyxl import load_workbook

from leadscraper import pipeline
from leadscraper.cache import Cache
from leadscraper.crawler import SiteCrawler
from leadscraper.excel import write_workbook
from leadscraper.funding import funding_reference_rows
from leadscraper.models import Company, Enrichment, Lead, SearchSpec
from leadscraper.places import PlacesClient

web = pytest.mark.usefixtures("fixture_web")


def _enrich(settings, website: str, name: str = "Test", **company_kw):
    async def go():
        crawler = SiteCrawler(settings)
        try:
            return await pipeline.enrich_company(
                Company(place_id="t", name=name, website=website, **company_kw), crawler
            )
        finally:
            await crawler.close()

    return asyncio.run(go())


@web
def test_makler_gf_mobile_found_outside_impressum(fast_settings):
    enr = _enrich(fast_settings, "https://www.rheinblick-immobilien-koeln.de/", user_rating_count=120)
    assert enr.legal_name == "Rheinblick Immobilien GmbH"
    assert enr.rechtsform == "GmbH" and enr.handelsregister == "HRB 87654" and enr.ustid == "DE298765432"
    assert (enr.impressum_street, enr.impressum_plz, enr.impressum_city) == (
        "Rheinuferstraße 12",
        "50667",
        "Köln",
    )
    gf = next(p for p in enr.people if p.name == "Thomas Berger")
    assert gf.role_category == "geschaeftsfuehrung"
    assert gf.mobile is not None and gf.mobile.e164 == "+491715550123"
    assert gf.mobile.source_url and "impressum" not in gf.mobile.source_url
    assert gf.email == "t.berger@rheinblick-immobilien-koeln.de"
    kranz = next(p for p in enr.people if p.name == "Julia Kranz")
    assert kranz.mobile is not None and kranz.mobile.e164 == "+491725550456"
    assert not any(p.name == "Lena Hoffmann" for p in enr.people)  # Web-Agentur im Impressum ignoriert
    assert "+492215550001" not in {p.e164 for p in enr.phones}  # Fax
    assert enr.size.point_estimate == 9 and enr.size.confidence == "high"
    assert enr.whatsapp_url == "https://wa.me/491715550123"
    assert enr.instagram_url and "instagram.com" in enr.instagram_url
    assert any("100 % Lehrgangskosten" in i for i in enr.call_indicators)


@web
def test_makler2_inhaberin_mobile_on_kontakt_and_whatsapp(fast_settings):
    enr = _enrich(fast_settings, "www.sonnenhof-immobilien-bonn.de")
    assert enr.rechtsform == "e.K."
    inh = next(p for p in enr.people if p.name == "Claudia Sonnenhof")
    assert inh.role_category == "inhaber"
    assert inh.mobile is not None and inh.mobile.e164 == "+491605550777"
    wa = [p for p in enr.phones if p.e164 == "+491605550777"]
    assert len(wa) == 1 and wa[0].kind == "mobile"
    brandt = next(p for p in enr.people if p.name == "Tim Brandt")
    assert brandt.mobile is not None and brandt.mobile.e164 == "+4915155507788"
    assert enr.size.point_estimate == 6


@web
def test_steuer_partner_mobile_via_tel_link_and_hr_contact(fast_settings):
    enr = _enrich(fast_settings, "https://www.weber-lind-steuerberater.de")
    assert enr.rechtsform == "PartG mbB"
    names = {p.name: p for p in enr.people}
    assert names["Markus Weber"].role_category == "geschaeftsfuehrung"
    assert names["Sonja Lind"].role_category == "geschaeftsfuehrung"
    assert names["Sonja Lind"].mobile is not None and names["Sonja Lind"].mobile.e164 == "+4915155509991"
    assert names["Anja Roth"].role_category == "hr" and names["Anja Roth"].email == "a.roth@weber-lind.de"
    assert enr.size.point_estimate == 23
    assert "HR-/Ausbildungsverantwortliche benannt" in enr.call_indicators
    assert "sucht Personal" in enr.call_indicators
    assert "KI/Digitalisierung auf der Website erwähnt" in enr.call_indicators


@web
def test_agentur_spa_no_mobile_and_no_fax(fast_settings):
    enr = _enrich(fast_settings, "https://www.pixelwerk-digital-agentur.de")
    assert {p.name for p in enr.people} == {"Lena Hoffmann", "Daniel Schuster"}
    assert all(p.role_category == "geschaeftsfuehrung" for p in enr.people)
    assert enr.mobiles == []
    assert enr.phones == []  # Fax verworfen, sonst nichts
    assert "hallo@pixelwerk-digital-agentur.de" in enr.emails


def test_no_website_uses_places_phone(fast_settings):
    company = Company(place_id="x", name="Makler Müller", phone="0170 5550999", user_rating_count=3)

    async def go():
        crawler = SiteCrawler(fast_settings)
        try:
            return await pipeline.enrich_company(company, crawler)
        finally:
            await crawler.close()

    enr = asyncio.run(go())
    assert enr.mobiles and enr.mobiles[0].source == "places"
    assert enr.size.confidence == "low"


@web
def test_search_dedupes_and_filters_chains(fast_settings):
    spec = SearchSpec(queries=["Immobilienmakler"], city="Köln", max_results_per_query=10)

    async def go():
        client = PlacesClient(fast_settings.google_places_api_key)
        try:
            return await pipeline.search_companies(spec, client)
        finally:
            await client.close()

    companies = asyncio.run(go())
    ids = [c.place_id for c in companies]
    assert ids == ["ChIJ-rheinblick", "ChIJ-sonnenhof", "ChIJ-weber", "ChIJ-pixelwerk", "ChIJ-nowebsite"]
    assert spec.lat is not None  # geocodiert
    assert companies[0].bundesland == "Nordrhein-Westfalen" and companies[0].plz == "50667"


@web
def test_end_to_end_build_leads_and_excel(fast_settings, tmp_path: Path):
    spec = SearchSpec(queries=["Immobilienmakler"], city="Köln", min_employees=5, max_employees=50)

    async def go():
        client = PlacesClient(fast_settings.google_places_api_key)
        try:
            companies = await pipeline.search_companies(spec, client)
        finally:
            await client.close()
        cache = Cache(fast_settings.cache_path)
        try:
            return await pipeline.build_leads(companies, spec, fast_settings, cache)
        finally:
            cache.close()

    leads = asyncio.run(go())
    names = [ld.display_name for ld in leads]
    # Die drei Firmen mit Entscheider-Handy stehen oben, SPA-Agentur und Firma ohne Website unten
    assert set(names[:3]) == {
        "Rheinblick Immobilien GmbH",
        "Sonnenhof Immobilien e.K.",
        "Weber & Lind Steuerberater PartG mbB",
    }
    assert all(ld.best_contact and ld.best_contact.mobile for ld in leads[:3])
    assert names[3:] == ["Pixelwerk Digital GmbH", "Makler Müller"]
    assert [ld.score for ld in leads] == sorted((ld.score for ld in leads), reverse=True)
    rb = next(ld for ld in leads if ld.display_name == "Rheinblick Immobilien GmbH")
    assert rb.best_contact.name == "Thomas Berger" and rb.best_mobile.e164 == "+491715550123"
    assert rb.funding.size_band == "<50" and rb.funding.lehrgangskosten_pct == 100
    assert rb.funding.landesprogramm == "Bildungsscheck NRW 2.0"
    assert rb.in_target_size is True
    px = next(ld for ld in leads if ld.display_name == "Pixelwerk Digital GmbH")
    assert px.best_mobile is None

    mobile_only = [
        ld for ld in leads if pipeline.passes_filters(ld, SearchSpec(queries=["x"], require_mobile=True))
    ]
    assert {ld.display_name for ld in mobile_only} >= {
        "Rheinblick Immobilien GmbH",
        "Sonnenhof Immobilien e.K.",
    }
    assert all(ld.best_mobile for ld in mobile_only)

    out = write_workbook(leads, spec, tmp_path / "leads.xlsx", funding_rows=funding_reference_rows())
    wb = load_workbook(out)
    ws = wb["Leads"]
    headers = [c.value for c in ws[1]]
    row = dict(zip(headers, [c.value for c in ws[2]], strict=False))
    assert row["Unternehmensname"] == leads[0].display_name
    assert row["Nummer"] == leads[0].best_contact.mobile.national
    assert row["Firmenname Quelle"] == "Impressum"
    assert row["Straße"] == leads[0].address[0]
    assert wb["Entscheider"].max_row >= 4


@web
def test_multiple_cities_are_searched_and_deduped(fast_settings, fixture_web):
    spec = SearchSpec(
        queries=["Immobilienmakler"], city="Köln", cities=["Bonn", "Köln"], max_results_per_query=10
    )

    async def go():
        client = PlacesClient(fast_settings.google_places_api_key)
        try:
            return await pipeline.search_companies(spec, client)
        finally:
            await client.close()

    companies = asyncio.run(go())
    searches = [r for r in fixture_web if "places:searchText" in r]
    assert len(searches) >= 4  # 2 Orte × (Geocode + Suche), Köln nur einmal
    assert len({c.place_id for c in companies}) == len(companies)


def test_warn_unreachable_reports_network_problem():
    messages: list[str] = []
    companies = [Company(place_id=str(i), name=f"F{i}", website=f"https://f{i}.de") for i in range(6)]
    enrichments = [
        Enrichment(website=c.website, errors=["Startseite nicht erreichbar (ConnectError: 403)"])
        for c in companies
    ]
    pipeline.warn_unreachable(companies, enrichments, messages.append)
    assert messages and "Netzwerk/Proxy" in messages[0] and "ConnectError" in messages[0]
    enrichments[0].pages_crawled = ["https://f0.de/"]
    enrichments[1].pages_crawled = ["https://f1.de/"]
    messages.clear()
    pipeline.warn_unreachable(companies, enrichments, messages.append)
    assert not messages  # nur 4 von 6 → keine Warnung


def test_employment_signal_on_fixture_sites(fast_settings, fixture_web):
    steuer = _enrich(fast_settings, "https://www.weber-lind-steuerberater.de")
    assert steuer.employment_signal == "angestellt"  # "Steuerfachangestellte", Azubi-Rollen
    assert any(e.startswith("+") for e in steuer.employment_evidence)
    makler = _enrich(fast_settings, "https://www.rheinblick-immobilien-koeln.de/")
    assert makler.employment_signal == "angestellt"  # "9 Mitarbeitern", Auszubildende, Büroleitung
    assert "Festangestellte erkennbar (förderfähig)" in makler.call_indicators


def test_freelance_signal_blocks_premium():
    from leadscraper.crawler import CrawlResult, Page
    from leadscraper.scoring import premium_check

    text = "Wir suchen freie Handelsvertreter auf Provisionsbasis. Ihr Ansprechpartner: Max Muster"
    page = Page(
        url="https://f.de/",
        final_url="https://f.de/",
        kind="startseite",
        html="",
        text=text,
        lines=[text],
        links=[],
    )
    crawl = CrawlResult(website="https://f.de/", pages=[page])
    enr = pipeline.build_enrichment(Company(place_id="f", name="F Immobilien", website="https://f.de"), crawl)
    assert enr.employment_signal == "frei"
    lead = Lead(company=Company(place_id="f", name="F Immobilien", website="https://f.de"), enrichment=enr)
    assert any("freie Handelsvertreter" in m for m in premium_check(lead, SearchSpec(queries=["x"])))


def test_run_cities_is_resumable(fast_settings, fixture_web, tmp_path: Path):
    spec = SearchSpec(queries=["Immobilienmakler"], max_results_per_query=10, radius_km=20)
    orte = [
        {"name": "Köln", "bundesland": "Nordrhein-Westfalen"},
        {"name": "Bonn", "bundesland": "Nordrhein-Westfalen", "radius_km": 6},
    ]
    jsonl = tmp_path / "de.jsonl"
    leads = asyncio.run(pipeline.run_cities(spec, fast_settings, orte, jsonl))
    assert leads and jsonl.exists() and jsonl.with_suffix(".state.json").exists()
    n_requests = len(fixture_web)
    again = asyncio.run(pipeline.run_cities(spec, fast_settings, orte, jsonl))
    assert len(fixture_web) == n_requests  # beide Orte erledigt → keine neuen Requests
    assert [ld.company.place_id for ld in again] == [ld.company.place_id for ld in leads]
    assert len({ld.company.place_id for ld in leads}) == len(leads)  # keine Duplikate über Orte hinweg


def test_load_orte_filter():
    orte = pipeline.load_orte(bundeslaender=["NRW", "Bremen"])
    assert orte and {o["bundesland"] for o in orte} == {"Nordrhein-Westfalen", "Bremen"}
    assert any(o["name"] == "Köln" for o in orte)
    assert len(pipeline.load_orte()) > 400


def test_read_leads_jsonl_skips_truncated_last_line(tmp_path: Path):
    """Ein abgebrochener Lauf hinterlässt oft eine halbe Zeile – der Neustart muss trotzdem fortsetzen."""
    from leadscraper.demo import demo_leads

    leads = demo_leads(SearchSpec(queries=["x"]))
    path = tmp_path / "de.jsonl"
    body = "\n".join(ld.model_dump_json() for ld in leads)
    path.write_text(body + "\n" + leads[0].model_dump_json()[:200], encoding="utf-8")
    assert len(pipeline.read_leads_jsonl(path)) == len(leads)


def test_refresh_lead_applies_indicators_without_crawl():
    """Gespeicherte Leads werden mit neuen Regeln bewertet – ohne Netz, ohne Google-Anfragen."""
    from leadscraper import funding
    from leadscraper.models import Person, PhoneNumber

    url = "https://b.de/team"
    gf = Person(name="Thomas Berger", role_category="geschaeftsfuehrung", source_url=url)
    staff = [
        Person(name=n, role="Immobilienmakler", source_url=url)
        for n in ("Julia Kranz", "Anna Heck", "Nina Lenz")
    ]
    mobile = PhoneNumber(
        raw="0171 5550123", e164="+491715550123", national="0171 5550123", kind="mobile", source="kontakt"
    )
    lead = Lead(
        company=Company(place_id="p", name="Berger Immobilien", website="https://b.de"),
        enrichment=Enrichment(
            pages_crawled=[url], people=[gf, *staff], phones=[mobile], emails=["t.berger@b.de"]
        ),
    )
    spec = SearchSpec(queries=["x"], min_employees=5, max_employees=50)
    out = pipeline.refresh_lead(lead, spec, funding.load_funding_config())
    assert out.enrichment.size.confidence == "medium"
    assert out.enrichment.size.employees_min == 4  # vier namentliche Personen belegen mindestens vier
    assert out.enrichment.employment_signal == "angestellt"
    # einzige Handynummer + genau ein Entscheider → eindeutig diesem zugeordnet
    assert out.enrichment.mobile_assignment == "eindeutig"
    assert out.enrichment.people[0].mobile is not None
    # vier belegte Köpfe reichen für die Mindestgröße 5 nicht
    assert out.enrichment.staff.headcount == 4
    assert out.premium is False and out.premium_missing == [
        "nur 4 Beschäftigte belegt (mindestens 5 gefordert)"
    ]

    # mit einer fünften Person auf der Team-Seite ist die Belegschaft belegt
    lead.enrichment.people.append(Person(name="Tim Brandt", role="Innendienst", source_url=url))
    out2 = pipeline.refresh_lead(lead, spec, funding.load_funding_config())
    assert out2.enrichment.staff.headcount == 5
    assert out2.premium is True and out2.premium_missing == []


def test_refresh_lead_discards_rating_as_headcount():
    """Alte Läufe haben „4,5/5 Mitarbeiter Zufriedenheit“ als Größe gespeichert – refresh verwirft das."""
    from leadscraper import funding
    from leadscraper.models import SizeEstimate

    lead = Lead(
        company=Company(place_id="p", name="X GmbH", website="https://x.de"),
        enrichment=Enrichment(
            pages_crawled=["https://x.de/"],
            size=SizeEstimate(
                point_estimate=5,
                employees_min=5,
                employees_max=5,
                confidence="high",
                evidence=["Text: „Kununu.com 4,5/5 Mitarbeiter Zufriedenheit“ (https://x.de/)"],
            ),
        ),
    )
    out = pipeline.refresh_lead(lead, SearchSpec(queries=["x"]), funding.load_funding_config())
    assert out.enrichment.size.point_estimate is None
    assert out.premium is False


def test_responsible_person_with_company_name_counts_as_owner():
    """Kleine Büros nennen im Impressum nur den Verantwortlichen nach § 18 MStV."""
    from leadscraper.models import Person

    enr = Enrichment(
        pages_crawled=["https://imzh.de/impressum/"],
        legal_name="Immobilienzentrum Hoffmann",
        people=[
            Person(name="Patrick Hoffmann", role="Inhaltlich verantwortlich", role_category="sonstige"),
            Person(name="Lea Wagner", role="Inhaltlich verantwortlich", role_category="sonstige"),
        ],
    )
    pipeline.promote_responsible_owner(enr, "IMZH Immobilienzentrum Hoffmann")
    assert enr.people[0].role_category == "inhaber"
    assert enr.people[1].role_category == "sonstige"  # fremder Nachname bleibt unverändert

    # Steht bereits ein Geschäftsführer im Impressum, wird nichts befördert
    enr2 = Enrichment(
        pages_crawled=["https://x.de/"],
        people=[
            Person(name="Anna Berger", role="Geschäftsführerin", role_category="geschaeftsfuehrung"),
            Person(name="Tim Berger", role="Verantwortlich", role_category="sonstige"),
        ],
    )
    pipeline.promote_responsible_owner(enr2, "Berger Immobilien GmbH")
    assert enr2.people[1].role_category == "sonstige"


def test_company_number_under_every_portrait_is_not_personal():
    """Wiederholt eine Team-Seite unter jedem Porträt dieselbe Nummer, gehört sie der Firma."""
    from leadscraper.models import PhoneNumber
    from leadscraper.pipeline import _drop_shared_numbers

    def pn(person):
        return PhoneNumber(
            raw="0178 3364585",
            e164="+491783364585",
            national="0178 3364585",
            kind="mobile",
            source="team",
            person=person,
        )

    alle = [pn("Helmut Jentz"), pn("Lisbeth Jentz"), pn("Sarah Jentz")]
    deduped = _drop_shared_numbers([pn("Helmut Jentz")], alle)
    assert deduped[0].person is None

    nur_einer = [pn("Helmut Jentz"), pn("Helmut Jentz")]
    assert _drop_shared_numbers([pn("Helmut Jentz")], nur_einer)[0].person == "Helmut Jentz"


def test_number_stays_with_the_person_across_pages():
    """Die Handynummer des Inhabers steht auf Team-, Objekt- und Kontaktseite – sie bleibt seine."""
    from leadscraper.models import PhoneNumber
    from leadscraper.pipeline import _drop_shared_numbers

    def pn(url, person=None):
        return PhoneNumber(
            raw="0171 5550123",
            e164="+491715550123",
            national="0171 5550123",
            kind="mobile",
            source="team",
            source_url=url,
            person=person,
        )

    ueberall = [pn("https://x.de/team/", "Thomas Berger"), pn("https://x.de/objekt/1", "Thomas Berger")]
    assert _drop_shared_numbers([pn("https://x.de/team/", "Thomas Berger")], ueberall)[0].person


def test_location_portal_counts_only_its_own_branch():
    """Franchise-Portale führen jeden Standort unter eigenem Pfad – fremde Standorte zählen nicht mit."""
    from leadscraper.pipeline import _scope_to_own_location

    urls = {
        "https://portal.de/makler-in-koeln/sued/",
        "https://portal.de/makler-in-koeln/sued/team/",
        "https://portal.de/makler-in-bochum/",
        "https://portal.de/team/",
    }
    eigene = _scope_to_own_location(urls, "https://portal.de/makler-in-koeln/sued/")
    assert eigene == {
        "https://portal.de/makler-in-koeln/sued/",
        "https://portal.de/makler-in-koeln/sued/team/",
    }
    # Eine gewöhnliche Firmenseite bleibt unangetastet
    normal = {"https://makler.de/team/", "https://makler.de/kontakt/"}
    assert _scope_to_own_location(normal, "https://makler.de/") == normal


def test_refresh_drops_evidence_that_no_longer_holds():
    """Verschärfte Textregeln wirken auch auf gespeicherte Läufe, ohne jede Website neu zu laden."""
    from leadscraper import funding
    from leadscraper.models import SizeEstimate

    def lead_mit(beleg: str) -> Lead:
        return Lead(
            company=Company(place_id="p", name="X GmbH", website="https://x.de"),
            enrichment=Enrichment(
                pages_crawled=["https://x.de/"],
                size=SizeEstimate(
                    point_estimate=17,
                    employees_min=17,
                    employees_max=17,
                    confidence="high",
                    evidence=[beleg],
                ),
            ),
        )

    spec = SearchSpec(queries=["x"], min_employees=5, max_employees=50)
    cfg = funding.load_funding_config()
    veraltet = pipeline.refresh_lead(
        lead_mit("Text: „1.271 Urteile wurden für die 17 Makler berücksichtigt“ (https://x.de/)"), spec, cfg
    )
    assert veraltet.enrichment.staff.headcount == 0
    gueltig = pipeline.refresh_lead(
        lead_mit("Text: „Unser Team aus 17 Mitarbeitern betreut Sie“ (https://x.de/)"), spec, cfg
    )
    assert gueltig.enrichment.staff.headcount == 17


def test_license_and_franchise_wording_marks_freelancers():
    """Betz Immobilien schreibt es selbst: ein Zusammenschluss rechtlich selbstständiger Makler."""
    from leadscraper.pipeline import _FREELANCE_RE

    for satz in (
        "Betz Immobilien ist ein Zusammenschluss von rechtlich selbstständigen Maklern.",
        "Jeder Lizenz- oder Franchisenehmer ist ein rechtlich eigenständiges Unternehmen.",
        "Sie sind Immobilienmakler und möchten auf selbstständiger Basis arbeiten?",
    ):
        assert _FREELANCE_RE.search(satz), satz
    assert not _FREELANCE_RE.search("Unsere Mitarbeiterin Mandy Schwarz berät Sie in Festanstellung.")


def test_haengende_website_blockiert_den_lauf_nicht(monkeypatch):
    """Ein Lauf über 1.600 Firmen stand anderthalb Stunden still, weil eine Verbindung nie zurückkam."""
    import asyncio

    from leadscraper import pipeline as pipeline_mod
    from leadscraper.models import Company
    from leadscraper.settings import Settings

    async def haengt(company, crawler):
        if "langsam" in (company.website or ""):
            await asyncio.sleep(30)
        return pipeline_mod.Enrichment(website=company.website, pages_crawled=["x"])

    monkeypatch.setattr(pipeline_mod, "enrich_company", haengt)
    firmen = [
        Company(place_id="1", name="Langsam GmbH", website="https://langsam.example"),
        Company(place_id="2", name="Schnell GmbH", website="https://schnell.example"),
    ]
    einstellungen = Settings(concurrency=2, cache_path=None, site_timeout_seconds=0.2)
    ergebnis = asyncio.run(pipeline_mod.enrich_all(firmen, einstellungen, None))
    assert ergebnis[0] is not None and "Zeitüberschreitung" in ergebnis[0].errors[0]
    assert ergebnis[1] is not None and ergebnis[1].pages_crawled == ["x"]
